"""
Utility script to export markdown documents to Microsoft Word (.docx) format.
Preserves headings, bullet points, bold styling, and tables.
"""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def md_to_docx(md_path: str, docx_path: str):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    doc = Document()
    
    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    in_table = False
    table_lines = []
    
    for line in lines:
        stripped = line.strip()
        
        # Check for table
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            continue
        elif in_table:
            # End of table, process it
            render_table(doc, table_lines)
            in_table = False
            table_lines = []
            
        if not stripped:
            continue
            
        # Headings
        if stripped.startswith('# '):
            p = doc.add_heading(stripped[2:], level=1)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
        elif stripped.startswith('## '):
            p = doc.add_heading(stripped[3:], level=2)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
        elif stripped.startswith('### '):
            p = doc.add_heading(stripped[4:], level=3)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
        elif stripped.startswith('#### '):
            p = doc.add_heading(stripped[5:], level=4)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            add_formatted_text(p, stripped[2:])
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            content = re.sub(r'^\d+\.\s', '', stripped)
            add_formatted_text(p, content)
        elif stripped.startswith('---'):
            # Horizontal rule
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, stripped)
            
    if in_table and table_lines:
        render_table(doc, table_lines)
        
    doc.save(docx_path)
    print(f"Exported {md_path} -> {docx_path}")


def add_formatted_text(paragraph, text: str):
    """Parses basic markdown bold and italics."""
    # Split by bold (**text**)
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def render_table(doc, table_lines: list[str]):
    """Parses markdown table lines and adds a styled Word table."""
    rows = []
    for line in table_lines:
        # Skip divider row (| --- | --- |)
        if re.match(r'^\|[\s\-:]+\|$', line.replace('|', '|').strip()):
            continue
        cells = [c.strip() for c in line.split('|')[1:-1]]
        if cells and not all(c.startswith('---') for c in cells):
            rows.append(cells)
            
    if not rows:
        return
        
    num_cols = max(len(r) for r in rows)
    num_rows = len(rows)
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Shading Accent 1' if 'Light Shading Accent 1' in [s.name for s in doc.styles] else 'Table Grid'
    
    for r_idx, row in enumerate(rows):
        for c_idx in range(min(len(row), num_cols)):
            cell = table.cell(r_idx, c_idx)
            cell.text = row[c_idx]
            if r_idx == 0:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True

if __name__ == '__main__':
    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    paper_dir = os.path.join(base, "paper")
    
    md_to_docx(
        os.path.join(paper_dir, "research_paper.md"),
        os.path.join(paper_dir, "research_paper.docx")
    )
    md_to_docx(
        os.path.join(paper_dir, "executive_summary.md"),
        os.path.join(paper_dir, "executive_summary.docx")
    )
